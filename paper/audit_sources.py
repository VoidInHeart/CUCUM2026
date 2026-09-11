"""只读核对正式结果，所有新证据写入paper/review；不调用正式导出入口。"""
from pathlib import Path
import sys, json, hashlib, dataclasses
import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
OUT = ROOT / 'paper' / 'review'
OUT.mkdir(parents=True, exist_ok=True)
def read(p): return json.loads((ROOT / p).read_text(encoding='utf-8'))
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def default(x):
    if isinstance(x,np.ndarray): return x.tolist()
    if isinstance(x,np.generic): return x.item()
    if hasattr(x,'isoformat'): return x.isoformat()
    raise TypeError(type(x).__name__)
def save(name,x): (OUT/name).write_text(json.dumps(x,ensure_ascii=False,indent=2,default=default),encoding='utf-8')
protected = [p for base in [ROOT/'src',ROOT/'data',ROOT/'target',ROOT.parent/'article'] for p in base.rglob('*') if p.is_file() and '__pycache__' not in p.parts]
before={str(p):sha(p) for p in protected}
save('source_hashes.json',before)
from src.question1.data_loader import load_input
from src.question1.optimizer import solve as solve1
from src.question1.validator import validate_solution
from src.question1.summary import build_table1_summary,build_table2_summary
from src.question2.data_loader import load_price,load_actual
from src.question2.pipeline import run_annual
from src.question2.summary import build_summaries
from src.question2.exporter import read_template_labels
from src.question3.storage import load_daily_results
from src.question3.validator import validate_complete
from src.question3.forecast_loader import load_forecasts
from src.question3.rolling_scenario_builder import RollingScenarioBuilder
from src.question3.rolling_controller import solve_day
from src.question3.evaluator import evaluate_actual_day
from src.question4.price_loader import load_annual_price
from src.question4.storage import load_analysis

report={}
data=load_input(); sol=solve1(data); validate_solution(data,sol)
w=load_workbook(ROOT/'target/附件5/result1.xlsx',read_only=True,data_only=True)
rows=list(w['计划购电量'].values); labels=[r[0] for r in rows[1:]]; saved=np.array([r[1] for r in rows[1:]])
report['q1']={'fresh_cost':sol.total_purchase_cost_yuan,'official_cost':float(data.price@saved),'grid_max_difference':float(np.max(abs(saved-sol.grid_purchase_kwh)))}
assert abs(report['q1']['fresh_cost']-report['q1']['official_cost'])<1e-5
q1={'table1':build_table1_summary(labels,saved,data.price),'table2':{},'fresh':dataclasses.asdict(sol),'baseline_cost':float(data.price@np.maximum(data.load_kwh-data.pv_kwh,0)), 'baseline_kwh':float(np.maximum(data.load_kwh-data.pv_kwh,0).sum())}
br=list(w['充放电量'].values)
for r in br[1:7]: q1['table2'][r[0]]={'累计充电量':r[1],'累计放电量':r[2]}
q1['table2']['0:00 储电量']=br[1][4];q1['table2']['24:00 储电量']=br[2][4]
w.close();save('question1_verified.json',q1)
price=load_price();actual=load_actual(price=price);labels=read_template_labels()
report['time_mapping']={'attachment_first':str(data.raw_time[0]),'attachment_last':str(data.raw_time[-1]),'template_first':labels[0],'template_last':labels[-1],'paper_10_index_zero_based':labels.index('10:00-10:10'),'internal_interval_at_that_index':'09:50-10:00','status':'保留位置映射；模板区间标签与内部结束时刻约定相差10分钟，未擅自平移结果'}
print('Recomputing Q2 334 days...',flush=True)
q2=run_annual(price,actual);summary2=build_summaries(q2,labels)
orig=read('target/question2/summary.json')
report['q2']={'days':len(q2),'total_cost_difference':summary2['annual']['total_cost']-orig['annual']['total_cost']}
assert abs(report['q2']['total_cost_difference'])<1e-4
from unittest.mock import patch
from scipy.optimize import linprog as actual_linprog
from scipy.sparse import vstack, lil_matrix
from src.question2.optimizer import solve as solve2
from src.question2.scenario_builder import build_rolling_scenarios
from src.question2.evaluator import evaluate
wb=load_workbook(ROOT/'target/附件5/result2.xlsx',read_only=True,data_only=True)
gridrows=list(wb['计划购电量'].values);battrows=list(wb['充放电量'].values);wb.close()
degenerate=[]
for i,r in enumerate(q2):
    official=np.array(gridrows[i+1][1:145],float)
    delta=float(np.max(abs(official-r.plan.grid_kwh)))
    if delta<1e-5:continue
    def fixed_linprog(objective,**kwargs):
        bounds=list(kwargs['bounds']);bounds[:144]=[(float(x),float(x)) for x in official];kwargs['bounds']=bounds
        extra=lil_matrix((12,len(objective)));values=[]
        for k in range(6):
            extra[2*k,144+k*24:144+(k+1)*24]=1;extra[2*k+1,288+k*24:288+(k+1)*24]=1
            values.extend(battrows[1+6*i+k][2:4])
        kwargs['A_eq']=vstack([kwargs['A_eq'],extra.tocsr()]);kwargs['b_eq']=np.r_[kwargs['b_eq'],values]
        return actual_linprog(objective,**kwargs)
    with patch('src.question2.optimizer.linprog',side_effect=fixed_linprog):
        plan=solve2(price.price,build_rolling_scenarios(actual,i+31))
    delta_obj=plan.expected_objective-r.plan.expected_objective
    assert abs(delta_obj)<1e-5
    q2[i]=evaluate(price.price,plan,actual.net_load_kwh[i+31])
    degenerate.append({'date':str(plan.date),'unconstrained_grid_difference':delta,'fixed_official_objective_difference':delta_obj})
report['q2']['equivalent_optima']=degenerate
print('Q2 fixed official plan checks complete',len(degenerate),flush=True)
q3=load_daily_results(ROOT/'target/question3/daily_results.jsonl');validate_complete(price.price,q3,(0,6,12))
forecasts=load_forecasts();builder=RollingScenarioBuilder(actual,forecasts)
errors=[]
for r in q3:
    if str(r.schedule.date) in read('target/question3/summary.json')['paper_dates']:
        fresh=evaluate_actual_day(price.price,solve_day(r.schedule.date,price.price,builder),actual.net_load_kwh[actual.dates.index(r.schedule.date)])
        errors.append({'date':str(r.schedule.date),'cost_difference':fresh.total_cost-r.total_cost,'grid_max_difference':float(np.max(abs(fresh.schedule.executed_grid-r.schedule.executed_grid)))})
report['q3']={'validated_days':len(q3),'fresh_four_days':errors}
assert max(abs(x['cost_difference']) for x in errors)<1e-4
prices=load_annual_price(actual)
q42,*_=load_analysis(ROOT/'target/question4/question4_2','4-2',prices)
q43,*_=load_analysis(ROOT/'target/question4/question4_3','4-3',prices)
report['q4']={'q42_validated_days':len(q42),'q43_validated_days':len(q43),'price_min':float(prices.price_yuan_per_kwh.min()),'price_max':float(prices.price_yuan_per_kwh.max()),'q3_q43_policy':[0,6,12]}

# 全部逐时合同、四小时储能汇总、紧急区间和日费用均对官方Excel只读对照。
report['workbooks']={}
for name,results,rolling in [('result2',q2,False),('result3',q3,True),('result4-2',q42,False),('result4-3',q43,True)]:
    wb=load_workbook(ROOT/f'target/附件5/{name}.xlsx',read_only=True,data_only=True)
    plans=list(wb['计划购电量'].values);battery=list(wb['充放电量'].values)
    adjusted=list(wb['调整购电量'].values) if rolling else None
    errors={'plan':0.,'adjusted':0.,'cost':0.,'energy':0.,'battery':0.,'emergency':0.}
    from src.question2.emergency import merge_emergency_intervals
    expected_events=[]
    for i,r in enumerate(results):
        if rolling:
            s=r.schedule;g=s.initial_plan.grid_kwh;ex=s.executed_grid;c=s.executed_charge;d=s.executed_discharge;u=r.real_emergency;energy=r.actual_grid_energy_kwh;day=s.date
        else:
            g=r.plan.grid_kwh;ex=g;c=r.plan.charge_kwh;d=r.plan.discharge_kwh;u=r.emergency_kwh;energy=r.total_grid_purchase;day=r.plan.date
        row=plans[i+1]
        assert str(row[0])[:10]==str(day)
        errors['plan']=max(errors['plan'],float(np.max(abs(np.array(row[1:145],float)-g))))
        errors['energy']=max(errors['energy'],abs(row[145]-(r.initial_plan_kwh if rolling else energy)))
        errors['cost']=max(errors['cost'],abs(row[146]-(r.plan_cost if rolling else r.total_cost)))
        if rolling:
            errors['adjusted']=max(errors['adjusted'],float(np.max(abs(np.array(adjusted[i+1][1:145],float)-ex))))
            errors['energy']=max(errors['energy'],abs(adjusted[i+1][145]-energy))
            errors['cost']=max(errors['cost'],abs(adjusted[i+1][146]-r.total_cost))
        for k in range(6):
            b=battery[1+6*i+k];sl=slice(24*k,24*(k+1))
            errors['battery']=max(errors['battery'],abs(b[2]-float(c[sl].sum())),abs(b[3]-float(d[sl].sum())))
        expected_events.extend((str(day),e.label,e.amount_kwh) for e in merge_emergency_intervals(labels,u))
    observed=[];day=None
    for row in list(wb['紧急购电量'].values)[1:]:
        if row[0] is not None:day=str(row[0])[:10]
        if row[1] is not None:observed.append((day,row[1],row[2]))
    assert len(observed)==len(expected_events),(name,len(observed),len(expected_events))
    for a,b in zip(observed,expected_events):
        assert a[:2]==b[:2],(name,a,b)
        errors['emergency']=max(errors['emergency'],abs(a[2]-b[2]))
    assert max(errors.values())<1e-4,(name,errors)
    report['workbooks'][name]={'days':len(results),'max_errors':errors,'events':len(observed)};wb.close()

report['physical']={}
for key,results,rolling in [('Q2',q2,False),('Q3',q3,True),('Q4-2',q42,False),('Q4-3',q43,True)]:
    maxima={'balance':0.,'soc':0.,'daily_boundary':0.,'simultaneous_slots':0}
    for r in results:
        if rolling:s=r.schedule;g=s.executed_grid;c=s.executed_charge;d=s.executed_discharge;e=s.executed_soc_end;n=r.actual_net_load;u=r.real_emergency;w=r.real_surplus
        else:s=r.plan;g=s.grid_kwh;c=s.charge_kwh;d=s.discharge_kwh;e=s.soc_end_kwh;n=r.actual_net_load_kwh;u=r.emergency_kwh;w=r.surplus_kwh
        maxima['balance']=max(maxima['balance'],float(np.max(abs(g+d-c+u-w-n))))
        maxima['soc']=max(maxima['soc'],float(np.max(abs(e-np.r_[6000,e[:-1]]-.9*c+d/.9))))
        maxima['daily_boundary']=max(maxima['daily_boundary'],float(abs(e[-1]-6000)))
        maxima['simultaneous_slots']+=int(np.sum((c>1e-6)&(d>1e-6)))
    report['physical'][key]=maxima
report['preservation']={'files':len(before),'changed':[p for p,h in before.items() if sha(Path(p))!=h]}
assert not report['preservation']['changed']
save('audit_results.json',report)
print(json.dumps(report,ensure_ascii=False,indent=2),flush=True)
