"""动态价格官方输出、失败保护及可复核存档。"""
from dataclasses import replace
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase, main
from unittest.mock import patch
import shutil
import numpy as np
from openpyxl import load_workbook

from src.question2.data_loader import load_actual
from src.question2.evaluator import evaluate
from src.question2 import exporter as exporter2
from src.question3 import exporter as exporter3
from src.question3.forecast_loader import load_forecasts
from src.question3.rolling_scenario_builder import RollingScenarioBuilder
from src.question3.evaluator import evaluate_actual_day
from src.question4 import config as cfg, question4_2, question4_3
from src.question4.price_loader import load_annual_price, get_daily_price
from src.question4.comparison import daily_metrics, compare
from src.question4.storage import save_analysis, load_analysis


class Question4OutputTests(TestCase):
    @classmethod
    def setUpClass(cls):
        actual = load_actual()
        prices = load_annual_price(actual)
        day = cfg.PAPER_DATES[1]
        q = get_daily_price(prices, day)
        b42 = question4_2.solve_day(day, prices, actual)
        builder = RollingScenarioBuilder(
            actual, load_forecasts(), cfg.Q4_3_SCENARIO_METHOD
        )
        b43 = question4_3.solve_day(day, prices, actual, builder)
        # 以已验证日构造完整导出日历，价格每天缩放，检查导出是否按日期选价。
        matrix = np.array([q*(1+i/365) for i in range(365)])
        cls.prices = replace(prices, price_yuan_per_kwh=matrix)
        cls.r42, cls.r43 = [], []
        for target in cfg.OUTPUT_DATES:
            daily_q = get_daily_price(cls.prices, target)
            cls.r42.append(evaluate(daily_q, replace(b42.plan, date=target), b42.actual_net_load_kwh))
            revisions = tuple(replace(r, plan=replace(r.plan, date=target),
                              adjustment_cashflow_yuan=float(1.5*daily_q[r.start_slot:] @ r.up_adjust -
                                                             .5*daily_q[r.start_slot:] @ r.down_adjust)) for r in b43.schedule.revisions)
            schedule = replace(b43.schedule, initial_plan=replace(b43.schedule.initial_plan, date=target), revisions=revisions)
            cls.r43.append(evaluate_actual_day(daily_q, schedule, b43.actual_net_load, cfg.Q4_3_UPDATE_HOURS))

    def test_result42_template_structure_and_daily_costs(self):
        with TemporaryDirectory() as folder:
            path = Path(folder)/"result4-2.xlsx"
            shutil.copy2(cfg.RESULT42_XLSX, path)
            exporter2.export_result(self.prices.daily_prices(), self.r42, path)
            wb = load_workbook(path)
            self.assertEqual(wb.sheetnames, ["计划购电量", "充放电量", "紧急购电量"])
            self.assertEqual((wb.active.max_row, wb.active.max_column), (335,147))
            self.assertEqual(wb["充放电量"].max_row, 2005)
            for i, r in enumerate(self.r42, 2):
                np.testing.assert_allclose([wb.active.cell(i,c).value for c in range(2,146)], r.plan.grid_kwh)
                self.assertAlmostEqual(wb.active.cell(i,146).value, r.total_grid_purchase)
                self.assertAlmostEqual(wb.active.cell(i,147).value, r.total_cost)
            wb.close()

    def test_result43_template_structure_and_daily_costs(self):
        with TemporaryDirectory() as folder:
            path = Path(folder)/"result4-3.xlsx"
            shutil.copy2(cfg.RESULT43_XLSX, path)
            exporter3.export_result(self.prices.daily_prices(), self.r43, path, update_hours=cfg.Q4_3_UPDATE_HOURS)
            wb = load_workbook(path)
            self.assertEqual(wb.sheetnames, ["计划购电量", "调整购电量", "充放电量", "紧急购电量"])
            self.assertEqual(wb["充放电量"].max_row, 2005)
            for i, r in enumerate(self.r43, 2):
                for sheet, values, energy, cost in ((wb.worksheets[0], r.schedule.initial_plan.grid_kwh, r.initial_plan_kwh, r.plan_cost),
                                                    (wb.worksheets[1], r.schedule.executed_grid, r.actual_grid_energy_kwh, r.total_cost)):
                    self.assertEqual((sheet.max_row,sheet.max_column), (335,147))
                    np.testing.assert_allclose([sheet.cell(i,c).value for c in range(2,146)], values)
                    self.assertAlmostEqual(sheet.cell(i,146).value, energy)
                    self.assertAlmostEqual(sheet.cell(i,147).value, cost)
            wb.close()

    def test_incomplete_or_wrong_price_export_preserves_template(self):
        with TemporaryDirectory() as folder:
            for original, exporter, results, kwargs in ((cfg.RESULT42_XLSX, exporter2, self.r42, {}),
                                                       (cfg.RESULT43_XLSX, exporter3, self.r43, {"update_hours": cfg.Q4_3_UPDATE_HOURS})):
                path = Path(folder)/original.name
                shutil.copy2(original,path)
                before = path.read_bytes()
                with self.assertRaises(ValueError): exporter.export_result(self.prices.daily_prices(), results[:-1], path, **kwargs)
                with self.assertRaises(ValueError): exporter.export_result(get_daily_price(self.prices,cfg.OUTPUT_DATES[0]), results, path, **kwargs)
                self.assertEqual(path.read_bytes(),before)

    def test_analysis_roundtrip_comparison_and_price_snapshot(self):
        with TemporaryDirectory() as folder:
            for model, results, exporter, path in (("4-2", self.r42, exporter2, cfg.RESULT42_XLSX), ("4-3", self.r43, exporter3, cfg.RESULT43_XLSX)):
                directory = Path(folder)/model
                rows = daily_metrics(model, self.prices, results)
                comparison = compare(model,rows,rows)
                self.assertEqual(next(r for r in comparison["changes"] if r["metric"]=="total_cost")["relative_change"],0)
                save_analysis(directory,model,self.prices,results,exporter.read_template_labels(path),rows,rows,comparison)
                restored, dynamic, fixed, output = load_analysis(directory,model,self.prices)
                self.assertEqual(len(restored),334)
                self.assertEqual(output,comparison)
                self.assertEqual(dynamic,fixed)
                self.assertAlmostEqual(restored[-1].total_cost,results[-1].total_cost)
                with self.assertRaisesRegex(ValueError,"prices differ"):
                    load_analysis(directory,model,replace(self.prices,price_yuan_per_kwh=self.prices.price_yuan_per_kwh*1.01))

    def test_failed_second_model_stops_before_both_official_exports(self):
        from src.question4.main import main as run
        with patch("sys.argv",["run_question4.py","--no-plots"]), \
             patch("src.question4.main.question4_2.run_annual",return_value=self.r42), \
             patch("src.question4.main.fixed_price_baseline",return_value=(self.prices,self.r42)), \
             patch("src.question4.main.save_analysis"), \
             patch("src.question4.main.question4_3.run_annual",side_effect=RuntimeError("2025-06-21 12:00 failed")), \
             patch("src.question4.main.exporter2.export_result") as export2, \
             patch("src.question4.main.exporter3.export_result") as export3:
            self.assertEqual(run(),1)
            export2.assert_not_called()
            export3.assert_not_called()


if __name__ == "__main__":
    main()
