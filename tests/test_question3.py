"""问题3：预报时序、冻结执行、退款结算与模板输出回归。"""
from dataclasses import replace
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest import TestCase, main
from unittest.mock import patch
import shutil

import numpy as np
import pandas as pd

from src.question2.data_loader import load_price, load_actual
from src.question3 import config as cfg
from src.question3.forecast_loader import load_forecasts
from src.question3.forecast_interpolator import interpolate_hourly_pv_forecast
from src.question3.rolling_scenario_builder import RollingScenarioBuilder, validate_scenarios
from src.question3.optimizer import solve_initial_plan, solve_adjustment
from src.question3.rolling_controller import solve_day, solve_day_causal, validate_policy
from src.question3.evaluator import evaluate_actual_day
from src.question3.validator import validate_schedule, validate_day, validate_revision, validate_complete
from src.question3.analysis import forecast_accuracy, result_update_hours
from src.question3.exporter import export_result, read_template_labels
from src.question3.summary import build_summary
from src.question3.storage import save_daily_results, load_daily_results
from src.question2.forecast import forecast_load
from src.question2.emergency import merge_emergency_intervals
from openpyxl import load_workbook


class Question3Tests(TestCase):
    @classmethod
    def setUpClass(cls):
        cls.price = load_price()
        cls.annual = load_actual(price=cls.price)
        cls.forecasts = load_forecasts()
        cls.builder = RollingScenarioBuilder(cls.annual, cls.forecasts)
        cls.day = date(2025, 6, 21)
        cls.index = cls.builder.date_index[cls.day]
        cls.schedule = solve_day(cls.day, cls.price.price, cls.builder)
        cls.result = evaluate_actual_day(cls.price.price, cls.schedule, cls.annual.net_load_kwh[cls.index])

    def test_attachment3_shape_issue_sequence_and_lengths(self):
        self.assertEqual(tuple(self.forecasts), cfg.ANNUAL_DATES)
        self.assertEqual(sum(len(item.issues) for item in self.forecasts.values()), 1460)
        for item in self.forecasts.values():
            self.assertEqual(tuple(item.issues), (0, 6, 12, 18))
            for hour, issue in item.issues.items():
                self.assertEqual(hour, issue.issue_hour)
                self.assertEqual(issue.hourly_power_kw.shape, (24,))

    def test_forecast_loader_groups_and_rejects_bad_input(self):
        frame = pd.read_excel(cfg.FORECAST_XLSX)
        # 组内非首行日期即使是无意义文本也不得解析为日期。
        mixed = frame.copy()
        mixed.iloc[1:4, 0] = "non-standard displayed date"
        with patch("src.question3.forecast_loader.pd.read_excel", return_value=mixed):
            self.assertEqual(len(load_forecasts()), 365)
        for case in ("shape", "date", "issue", "nan", "negative"):
            damaged = frame.copy()
            if case == "shape": damaged = damaged.iloc[:-1]
            elif case == "date": damaged.iloc[0, 0] = "2025-01-02"
            elif case == "issue": damaged.iloc[2, 1] = "18:00"
            elif case == "nan": damaged.iloc[2, 3] = np.nan
            else: damaged.iloc[2, 3] = -1
            with self.subTest(case=case), patch("src.question3.forecast_loader.pd.read_excel", return_value=damaged):
                with self.assertRaises(ValueError): load_forecasts()

    def test_midpoint_interpolation_and_output_lengths(self):
        for hour in cfg.ISSUE_HOURS:
            values = interpolate_hourly_pv_forecast(np.arange(1, 25) * 600.0, hour, 0.0)
            self.assertEqual(values.shape, (6 * (24 - hour),))
            np.testing.assert_allclose(values, (np.arange(len(values)) + 0.5) * 100)
        anchored = interpolate_hourly_pv_forecast(np.full(24, 600.0), 6, 1200.0)
        self.assertAlmostEqual(anchored[0], 1150.0)
        with self.assertRaises(ValueError): interpolate_hourly_pv_forecast(np.zeros(23), 0, 0)

    def test_history_dates_and_scenario_shapes(self):
        for target in (cfg.START_DATE, self.day, cfg.END_DATE):
            for hour in cfg.ISSUE_HOURS:
                scenarios = self.builder.build(target, hour)
                validate_scenarios(scenarios)
                self.assertTrue(all(day < target for day in scenarios.history_dates))
                self.assertEqual(scenarios.net_load_kwh.shape, (31, 6 * (24 - hour)))
                self.assertAlmostEqual(scenarios.probability.sum(), 1)
                self.assertGreaterEqual(scenarios.pv_scenario_kw.min(), 0)

    def test_paired_residual_scenarios_keep_same_day_errors_and_dynamic_count(self):
        builder = RollingScenarioBuilder(self.annual, self.forecasts, "paired_residual_weighted")
        first = builder.build(cfg.START_DATE, 0)
        self.assertEqual(len(first.history_dates), 30)
        self.assertGreater(np.ptp(first.probability), 0)
        scenarios = builder.build(self.day, 12)
        validate_scenarios(scenarios)
        self.assertEqual(len(scenarios.history_dates), 31)
        history_index = builder.date_index[scenarios.history_dates[-1]]
        start = cfg.ISSUE_SLOT[12]
        expected_load = np.maximum(
            scenarios.load_forecast_kw + self.annual.load_kw[history_index, start:]
            - forecast_load(self.annual, history_index)[start:], 0
        )
        expected_pv = np.maximum(
            scenarios.forecast_kw + builder.errors.get(
                scenarios.history_dates[-1], 12, before=self.day
            ), 0
        )
        np.testing.assert_allclose(scenarios.load_scenario_kw[-1], expected_load)
        np.testing.assert_allclose(scenarios.pv_scenario_kw[-1], expected_pv)
        solve_initial_plan(self.price.price, first)

    def test_paired_residual_scenarios_do_not_read_target_future(self):
        hour = 12
        start = cfg.ISSUE_SLOT[hour]
        load, pv = self.annual.load_kw.copy(), self.annual.pv_kw.copy()
        load[self.index, start:] += 1e6
        pv[self.index, start:] += 1e6
        load[self.index + 1:] += 1e6
        pv[self.index + 1:] += 1e6
        altered = replace(self.annual, load_kw=load, pv_kw=pv)
        allowed = {key: value for key, value in self.forecasts[self.day].issues.items() if key <= hour}
        forecasts = {**self.forecasts, self.day: replace(self.forecasts[self.day], issues=allowed)}
        original = RollingScenarioBuilder(
            self.annual, self.forecasts, "paired_residual_weighted"
        ).build(self.day, hour)
        changed = RollingScenarioBuilder(
            altered, forecasts, "paired_residual_weighted"
        ).build(self.day, hour)
        np.testing.assert_allclose(original.net_load_kwh, changed.net_load_kwh, rtol=0, atol=1e-9)
        np.testing.assert_array_equal(original.probability, changed.probability)

    def test_future_actual_and_unreleased_forecasts_not_accessed(self):
        for hour in cfg.ISSUE_HOURS:
            start = cfg.ISSUE_SLOT[hour]
            load, pv = self.annual.load_kw.copy(), self.annual.pv_kw.copy()
            load[self.index, start:] += 1e6
            pv[self.index, start:] += 1e6
            load[self.index + 1:] += 1e6
            pv[self.index + 1:] += 1e6
            altered = replace(self.annual, load_kw=load, pv_kw=pv)
            allowed_issues = {key: value for key, value in self.forecasts[self.day].issues.items() if key <= hour}
            forecasts = {**self.forecasts, self.day: replace(self.forecasts[self.day], issues=allowed_issues)}
            restricted = RollingScenarioBuilder(altered, forecasts)
            original, changed = self.builder.build(self.day, hour), restricted.build(self.day, hour)
            np.testing.assert_array_equal(original.net_load_kwh, changed.net_load_kwh)
            if hour == 0:
                np.testing.assert_allclose(solve_initial_plan(self.price.price, original).grid_kwh,
                                           solve_initial_plan(self.price.price, changed).grid_kwh)
        with self.assertRaisesRegex(ValueError, "precede"):
            self.builder.errors.get(self.day, 6, before=self.day)

    def test_history_error_uses_same_observed_anchor(self):
        historical = self.builder.annual.dates[self.index - 1]
        for hour in cfg.ISSUE_HOURS:
            start = cfg.ISSUE_SLOT[hour]
            anchor = self.annual.pv_kw[self.index - 1, start - 1] if start else 0
            prediction = interpolate_hourly_pv_forecast(self.forecasts[historical].issues[hour].hourly_power_kw, hour, anchor)
            expected = self.annual.pv_kw[self.index - 1, start:] - prediction
            np.testing.assert_allclose(self.builder.errors.get(historical, hour, before=self.day), expected)

    def test_known_initial_and_adjustment_objectives(self):
        price = np.ones(144)
        source = self.builder.build(self.day, 0)
        initial = replace(
            source,
            load_scenario_kw=source.pv_scenario_kw + 600.0,
            net_load_kwh=np.full((31, 144), 100.0),
        )
        self.assertAlmostEqual(solve_initial_plan(price, initial).solver_objective, 14400, places=5)
        source = self.builder.build(self.day, 18)
        scenarios = replace(
            source,
            load_scenario_kw=source.pv_scenario_kw + 600.0,
            net_load_kwh=np.full((31, 36), 100.0),
        )
        down = solve_adjustment(price, scenarios, np.full(36, 200.0), 6000)
        up = solve_adjustment(price, scenarios, np.zeros(36), 6000)
        self.assertAlmostEqual(down.adjustment_cashflow_yuan, -1800, places=5)
        self.assertAlmostEqual(up.adjustment_cashflow_yuan, 5400, places=5)
        self.assertAlmostEqual(down.plan.solver_objective, -1800, places=5)
        self.assertTrue((down.down_adjust <= down.old_grid + cfg.TOL).all())
        validate_revision(price, down)
        clipped = solve_adjustment(
            price, scenarios, np.zeros(36), cfg.SOC_MAX_KWH + cfg.TOL / 2
        )
        self.assertEqual(clipped.plan.initial_soc_kwh, cfg.SOC_MAX_KWH)

    def test_boundary_soc_and_frozen_decisions(self):
        s = self.schedule
        validate_schedule(self.price.price, s, cfg.Q3_FINAL_UPDATE_HOURS)
        np.testing.assert_array_equal(s.executed_grid[:36], s.initial_plan.grid_kwh[:36])
        for i, revision in enumerate(s.revisions):
            start = revision.start_slot
            stop = s.revisions[i + 1].start_slot if i + 1 < len(s.revisions) else 144
            np.testing.assert_array_equal(s.executed_grid[start:stop], revision.new_grid[:stop-start])
            self.assertAlmostEqual(revision.plan.initial_soc_kwh, s.executed_soc_end[start - 1])
            expected_old = s.initial_plan.grid_kwh[start:] if i == 0 else s.revisions[i - 1].new_grid[36:]
            np.testing.assert_allclose(revision.old_grid, expected_old)
        self.assertAlmostEqual(s.executed_soc_end[-1], 6000)
        damaged = s.executed_grid.copy()
        damaged[0] += 1
        with self.assertRaisesRegex(ValueError, "contract revision"):
            validate_schedule(self.price.price, replace(s, executed_grid=damaged), cfg.Q3_FINAL_UPDATE_HOURS)
        with self.assertRaises(ValueError): s.executed_grid[0] = 0

    def test_actual_balance_and_cost_accounting(self):
        result = self.result
        s = result.schedule
        np.testing.assert_allclose(s.executed_grid + s.executed_discharge - s.executed_charge + result.real_emergency - result.real_surplus,
                                   result.actual_net_load)
        self.assertAlmostEqual(result.total_cost, result.plan_cost + sum(v.adjustment_cashflow_yuan for v in s.revisions) + result.emergency_cost)
        self.assertAlmostEqual(result.emergency_cost, 5 * self.price.price @ result.real_emergency)
        with self.assertRaisesRegex(ValueError, "accounting"):
            validate_day(self.price.price, replace(result, total_cost=result.total_cost + 1), cfg.Q3_FINAL_UPDATE_HOURS)

    def test_default_policy_uses_only_selected_forecasts(self):
        self.assertEqual(cfg.Q3_FINAL_UPDATE_HOURS, (0, 6, 12))
        with patch.object(self.builder, "build", wraps=self.builder.build) as build:
            schedule = solve_day(self.day, self.price.price, self.builder)
        self.assertEqual([call.args[1] for call in build.call_args_list], [0, 6, 12])
        self.assertEqual(tuple(r.issue_hour for r in schedule.revisions), (6, 12))
        np.testing.assert_array_equal(schedule.executed_grid[72:], schedule.revisions[-1].new_grid)
        historical = solve_day(self.day, self.price.price, self.builder, cfg.ISSUE_HOURS)
        with self.assertRaisesRegex(ValueError, "revision sequence"):
            evaluate_actual_day(self.price.price, historical, self.annual.net_load_kwh[self.index])
        result = evaluate_actual_day(self.price.price, historical, self.annual.net_load_kwh[self.index], cfg.ISSUE_HOURS)
        self.assertEqual(result_update_hours([result]), cfg.ISSUE_HOURS)
        with self.assertRaisesRegex(ValueError, "consistent"):
            result_update_hours([result, self.result])

    def test_ablation_policy_execution(self):
        for policy in cfg.POLICIES:
            s = solve_day(self.day, self.price.price, self.builder, policy)
            self.assertEqual(tuple(r.issue_hour for r in s.revisions), policy[1:])
            evaluate_actual_day(self.price.price, s, self.annual.net_load_kwh[self.index], policy)
        for policy in ((6,), (0, 12, 6), (0, 6, 6), (0, 3)):
            with self.assertRaises(ValueError): validate_policy(policy)

    def test_causal_storage_updates_each_slot_but_only_revises_selected_contracts(self):
        builder = RollingScenarioBuilder(self.annual, self.forecasts, cfg.SCENARIO_METHOD)
        actual = self.annual.net_load_kwh[self.index]
        with patch.object(builder, "build", wraps=builder.build) as build:
            schedule = solve_day_causal(
                self.day, self.price.price, builder, actual, cfg.Q3_FINAL_UPDATE_HOURS
            )
        self.assertEqual(schedule.storage_execution, "causal_mpc")
        self.assertEqual([call.args[1] for call in build.call_args_list], [0, 6, 12])
        self.assertEqual(tuple(r.issue_hour for r in schedule.revisions), (6, 12))
        for revision in schedule.revisions:
            self.assertAlmostEqual(
                revision.plan.initial_soc_kwh,
                schedule.executed_soc_end[revision.start_slot - 1],
                places=6,
            )
        result = evaluate_actual_day(
            self.price.price, schedule, actual, cfg.Q3_FINAL_UPDATE_HOURS
        )
        validate_day(self.price.price, result, cfg.Q3_FINAL_UPDATE_HOURS)
        self.assertAlmostEqual(schedule.executed_soc_end[-1], cfg.FINAL_SOC_KWH, places=6)
        with TemporaryDirectory() as directory:
            path = Path(directory) / "causal.jsonl"
            save_daily_results(path, [result])
            restored = load_daily_results(path)[0]
            self.assertEqual(restored.schedule.storage_execution, "causal_mpc")
            validate_day(self.price.price, restored, cfg.Q3_FINAL_UPDATE_HOURS)

    def test_causal_storage_prefix_does_not_see_future_actual_net_load(self):
        builder = RollingScenarioBuilder(self.annual, self.forecasts, cfg.SCENARIO_METHOD)
        actual = self.annual.net_load_kwh[self.index].copy()
        changed = actual.copy()
        cutoff = 47
        changed[cutoff + 1:] += 1e6
        first = solve_day_causal(self.day, self.price.price, builder, actual, (0,))
        second = solve_day_causal(self.day, self.price.price, builder, changed, (0,))
        np.testing.assert_array_equal(first.executed_grid, second.executed_grid)
        for field in ("executed_charge", "executed_discharge", "executed_soc_end"):
            np.testing.assert_allclose(
                getattr(first, field)[:cutoff + 1], getattr(second, field)[:cutoff + 1],
                rtol=0, atol=cfg.RESIDUAL_TOL,
            )

    def test_forecast_accuracy_samples(self):
        scores = forecast_accuracy(self.annual, self.forecasts)
        for score in scores:
            self.assertEqual(score["samples"], 334 * 6 * (24 - score["issue_hour"]))
            self.assertGreaterEqual(score["RMSE_kW"], score["MAE_kW"])
            self.assertEqual(score["common_18_24_samples"], 334 * 36)

    def test_solver_failure_context(self):
        with patch("src.question3.optimizer.linprog", return_value=SimpleNamespace(success=False, status=2, message="infeasible")):
            with self.assertRaisesRegex(RuntimeError, "2025-06-21 0:00.*status 2.*infeasible"):
                solve_initial_plan(self.price.price, self.builder.build(self.day, 0))


class Question3ExportTests(TestCase):
    @classmethod
    def setUpClass(cls):
        Question3Tests.setUpClass()
        cls.price = Question3Tests.price.price
        base = Question3Tests.result
        cls.results = []
        # 合成完整日历：复用一个物理可行日，不在导出单测里重复进行全年优化。
        # 完整真实334日及六策略由运行入口单独执行并校验。
        for day in cfg.OUTPUT_DATES:
            initial = replace(base.schedule.initial_plan, date=day)
            revisions = tuple(replace(r, plan=replace(r.plan, date=day)) for r in base.schedule.revisions)
            cls.results.append(replace(base, schedule=replace(base.schedule, initial_plan=initial, revisions=revisions)))
        cls.labels = read_template_labels()

    def test_result3_all_sheets_and_totals(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "result3.xlsx"
            shutil.copy2(cfg.RESULT_XLSX, path)
            before = load_workbook(path)
            headers = [tuple(cell.value for cell in before[name][1]) for name in before.sheetnames]
            before.close()
            export_result(self.price, self.results, path)
            workbook = load_workbook(path)
            self.assertEqual(workbook.sheetnames, ["计划购电量", "调整购电量", "充放电量", "紧急购电量"])
            for i, sheet in enumerate(workbook):
                self.assertEqual(tuple(cell.value for cell in sheet[1]), headers[i])
            initial, adjusted, battery, emergency = (workbook[name] for name in workbook.sheetnames)
            self.assertEqual((initial.max_row, initial.max_column), (335, 147))
            self.assertEqual((adjusted.max_row, adjusted.max_column), (335, 147))
            self.assertEqual((battery.max_row, battery.max_column), (2005, 6))
            expected_events = []
            for i, result in enumerate(self.results):
                s = result.schedule
                for sheet, energy, ep, eq in ((initial, s.initial_plan.grid_kwh, result.initial_plan_kwh, result.plan_cost),
                                              (adjusted, s.executed_grid, result.actual_grid_energy_kwh, result.total_cost)):
                    self.assertEqual(sheet.cell(i + 2, 1).value.date(), s.date)
                    np.testing.assert_allclose([sheet.cell(i + 2, c).value for c in range(2, 146)], energy)
                    self.assertAlmostEqual(sheet.cell(i + 2, 146).value, ep)
                    self.assertAlmostEqual(sheet.cell(i + 2, 147).value, eq)
                for block in range(6):
                    row = 2 + i * 6 + block
                    self.assertEqual(battery.cell(row, 1).value.date() if block == 0 else battery.cell(row, 1).value,
                                     s.date if block == 0 else None)
                    self.assertAlmostEqual(battery.cell(row, 3).value, s.executed_charge[block * 24:(block + 1) * 24].sum())
                    self.assertAlmostEqual(battery.cell(row, 4).value, s.executed_discharge[block * 24:(block + 1) * 24].sum())
                    self.assertEqual(battery.cell(row, 6).value, 6000 if block < 2 else None)
                for j, event in enumerate(merge_emergency_intervals(self.labels, result.real_emergency)):
                    expected_events.append((s.date if j == 0 else None, event.label, event.amount_kwh))
            rows = list(emergency.values)[1:]
            self.assertEqual(len(rows), len(expected_events))
            for actual, expected in zip(rows, expected_events):
                self.assertEqual(actual[0].date() if actual[0] else None, expected[0])
                self.assertEqual(actual[1], expected[1])
                self.assertAlmostEqual(actual[2], expected[2])
            workbook.close()

    def test_export_failure_preserves_official_file(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "result3.xlsx"
            shutil.copy2(cfg.RESULT_XLSX, path)
            original = path.read_bytes()
            with self.assertRaisesRegex(ValueError, "334"):
                export_result(self.price, self.results[:-1], path)
            self.assertEqual(path.read_bytes(), original)
            with patch("src.question3.exporter.os.replace", side_effect=PermissionError("locked")):
                with self.assertRaisesRegex(RuntimeError, "locked"):
                    export_result(self.price, self.results, path)
            self.assertEqual(path.read_bytes(), original)
            self.assertEqual(list(Path(directory).glob("*.tmp.xlsx")), [])
            workbook = load_workbook(path)
            workbook.active.title = "wrong"
            workbook.save(path)
            workbook.close()
            original = path.read_bytes()
            with self.assertRaisesRegex(RuntimeError, "four official sheets"):
                export_result(self.price, self.results, path)
            self.assertEqual(path.read_bytes(), original)

    def test_paper_summary_uses_executed_schedule(self):
        summary = build_summary(self.results, self.labels)
        self.assertEqual(summary["update_hours"], [0, 6, 12])
        self.assertEqual(summary["policy"], "0+6+12")
        self.assertEqual(summary["storage_execution"], "frozen")
        self.assertEqual(set(summary["paper_dates"]), {str(day) for day in cfg.PAPER_DATES})
        for day, tables in summary["paper_dates"].items():
            result = self.results[cfg.OUTPUT_DATES.index(date.fromisoformat(day))]
            for label, values in tables["表1"].items():
                slot = self.labels.index(label)
                self.assertEqual(values["执行合同购电量"], result.schedule.executed_grid[slot])
            self.assertAlmostEqual(tables["全天指标"]["total_cost"], result.total_cost)

    def test_main_selects_official_policy_for_summary_storage_and_export(self):
        from src.question3.main import main as run_main
        import contextlib
        import io
        main_results = self.results
        historical = [replace(result, total_cost=result.total_cost + 1) for result in self.results]
        def run_policy(hours, *args):
            return main_results if hours == cfg.Q3_FINAL_UPDATE_HOURS else historical
        with patch("sys.argv", ["run_question3.py", "--no-plots"]), \
             patch("src.question3.main.evaluate_update_policy", side_effect=run_policy) as solve, \
             patch("src.question3.main.summarize_policies", return_value=[]) as ablate, \
             patch("src.question3.main.evaluate_model_ablation", return_value=[]) as model_ablate, \
             patch("src.question3.main.forecast_accuracy", return_value=[]), \
             patch("src.question3.main.build_summary", return_value={"annual": {}}) as summary, \
             patch("src.question3.main.save_analysis") as save, \
             patch("src.question3.main.export_result") as export, contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(run_main(), 0)
            self.assertEqual(solve.call_args_list[0].args[0], (0,6,12))
            self.assertEqual({call.args[0] for call in solve.call_args_list}, set(cfg.POLICIES))
            self.assertIs(ablate.call_args.args[0][cfg.ISSUE_HOURS], historical)
            self.assertIs(model_ablate.call_args.args[3][cfg.Q3_FINAL_UPDATE_HOURS], main_results)
            self.assertIs(summary.call_args.args[0], main_results)
            self.assertIs(save.call_args.args[1], main_results)
            self.assertIs(export.call_args.args[1], main_results)
            self.assertEqual(export.call_args.kwargs["update_hours"], (0,6,12))

    def test_four_update_archive_cannot_replace_official_by_default(self):
        schedule = solve_day(Question3Tests.day, self.price, Question3Tests.builder, cfg.ISSUE_HOURS)
        archived = evaluate_actual_day(self.price, schedule, Question3Tests.annual.net_load_kwh[Question3Tests.index], cfg.ISSUE_HOURS)
        records = [replace(archived, schedule=replace(schedule,
                   initial_plan=replace(schedule.initial_plan, date=day),
                   revisions=tuple(replace(r, plan=replace(r.plan, date=day)) for r in schedule.revisions))) for day in cfg.OUTPUT_DATES]
        validate_complete(self.price, records, cfg.ISSUE_HOURS)
        with TemporaryDirectory() as directory:
            path = Path(directory)/"result3.xlsx"
            shutil.copy2(cfg.RESULT_XLSX, path)
            original = path.read_bytes()
            with self.assertRaisesRegex(ValueError, "revision sequence"):
                export_result(self.price, records, path)
            self.assertEqual(path.read_bytes(), original)

    def test_failed_policy_stops_main_before_export(self):
        from src.question3.main import main as run_main
        with patch("sys.argv", ["run_question3.py", "--no-plots"]), \
             patch("src.question3.main.evaluate_update_policy", side_effect=RuntimeError("2025-03-20 12:00 failed")), \
             patch("src.question3.main.export_result") as export:
            self.assertEqual(run_main(), 1)
            export.assert_not_called()


if __name__ == "__main__":
    main()
