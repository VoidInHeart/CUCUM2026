"""运行：python -m unittest discover -s tests -v。结果仅写入临时目录。"""
from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase, main
from unittest.mock import patch
import shutil

import numpy as np
from openpyxl import load_workbook

from src.question1 import config as cfg
from src.question1.data_loader import load_input
from src.question1.exporter import export_result
from src.question1.optimizer import solve
from src.question1.summary import build_table1_summary, build_table2_summary
from src.question1.validator import validate_solution


class Question1Tests(TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = load_input()
        cls.solution = solve(cls.data)

    def test_input_shape_and_order(self):
        self.assertEqual(self.data.price.shape, (144,))
        workbook = load_workbook(cfg.INPUT_XLSX)
        self.assertEqual(self.data.raw_time, [row[0] for row in list(workbook.active.values)[1:]])
        workbook.close()

    def test_power_to_energy_conversion(self):
        self.assertAlmostEqual(6000 * cfg.TIME_STEP_HOURS, 1000)
        np.testing.assert_allclose(self.data.load_kwh, self.data.load_kw / 6)
        np.testing.assert_allclose(self.data.pv_kwh, self.data.pv_kw / 6)

    def test_battery_slot_limit(self):
        self.assertAlmostEqual(cfg.BATTERY_ENERGY_MAX_PER_SLOT_KWH, 5000 / 6)

    def test_solution_feasible(self):
        report = validate_solution(self.data, self.solution)
        self.assertLessEqual(max(report.values()), 1e-5)

    def test_terminal_soc(self):
        self.assertAlmostEqual(self.solution.soc_end_kwh[-1], 6000, places=5)

    def test_reference_objective(self):
        self.assertAlmostEqual(self.solution.total_purchase_cost_yuan, 35126.9486, delta=0.1)
        self.assertAlmostEqual(self.solution.total_purchase_kwh, 59482.6990, delta=0.1)

    def test_invalid_input(self):
        with TemporaryDirectory() as folder:
            path = Path(folder) / "input.xlsx"
            for case in ("rows", "header", "nan", "text", "inf", "negative"):
                with self.subTest(case=case):
                    workbook = load_workbook(cfg.INPUT_XLSX)
                    sheet = workbook.active
                    if case == "rows":
                        sheet.delete_rows(145)
                    else:
                        address, value = {"header": ("B1", "wrong"), "nan": ("C2", None),
                                          "text": ("B2", "bad"), "inf": ("D2", "inf"),
                                          "negative": ("C2", -1)}[case]
                        sheet[address] = value
                    workbook.save(path)
                    workbook.close()
                    with self.assertRaisesRegex(ValueError, str(path.name)):
                        load_input(path)
            with self.assertRaises(FileNotFoundError):
                load_input(Path(folder) / "missing.xlsx")

    def test_validator_rejects_corruption(self):
        for field, value in (("grid_purchase_kwh", np.zeros(143)),
                             ("charge_kwh", np.full(144, np.nan)),
                             ("discharge_kwh", np.full(144, 1000.0)),
                             ("curtailment_kwh", np.full(144, -1.0)),
                             ("soc_end_kwh", np.full(144, 6001.0)),
                             ("total_purchase_cost_yuan", float("inf"))):
            with self.subTest(field=field):
                damaged = deepcopy(self.solution)
                setattr(damaged, field, value)
                with self.assertRaises(ValueError):
                    validate_solution(self.data, damaged)

    def test_solver_failure(self):
        from types import SimpleNamespace
        with patch("src.question1.optimizer.linprog", return_value=SimpleNamespace(success=False, message="infeasible")):
            with self.assertRaisesRegex(RuntimeError, "infeasible"):
                solve(self.data)

    def test_surplus_pv_and_zero_prices(self):
        data = deepcopy(self.data)
        data.price[:] = 0
        data.pv_kw[:] = 30000
        data.pv_kwh[:] = 5000
        solution = solve(data)
        validate_solution(data, solution)
        self.assertGreater(solution.curtailment_kwh.sum(), 0)
        self.assertAlmostEqual(solution.total_purchase_cost_yuan, 0)

    def test_output_template(self):
        with TemporaryDirectory() as folder:
            path = Path(folder) / "result1.xlsx"
            shutil.copy2(cfg.RESULT_XLSX, path)
            before = load_workbook(path)
            table1, table2 = export_result(self.data, self.solution, path)
            after = load_workbook(path)
            self.assertEqual(before.sheetnames, after.sheetnames)
            allowed = {"计划购电量": {f"B{r}" for r in range(2, 146)},
                       "充放电量": {f"{c}{r}" for r in range(2, 8) for c in "BC"} | {"E2", "E3"}}
            for sheet in before:
                result = after[sheet.title]
                self.assertEqual(sheet.max_row, result.max_row)
                self.assertEqual(sheet.max_column, result.max_column)
                self.assertEqual(str(sheet.merged_cells), str(result.merged_cells))
                for row in sheet:
                    for cell in row:
                        updated = result[cell.coordinate]
                        if cell.coordinate in allowed[sheet.title]:
                            self.assertIsInstance(updated.value, (int, float))
                            self.assertEqual(updated.number_format, "0.000000")
                        else:
                            self.assertEqual(cell.value, updated.value)
                            self.assertEqual(cell._style, updated._style)
            np.testing.assert_allclose([after["计划购电量"].cell(r, 2).value for r in range(2, 146)],
                                       self.solution.grid_purchase_kwh)
            self.assertAlmostEqual(sum(v["累计充电量"] for v in table2.values() if isinstance(v, dict)),
                                   self.solution.charge_kwh.sum())
            self.assertAlmostEqual(table1["全天购电费"], self.solution.total_purchase_cost_yuan)
            before.close()
            after.close()

    def test_export_failures_preserve_file(self):
        with TemporaryDirectory() as folder:
            path = Path(folder) / "result1.xlsx"
            with self.assertRaises(FileNotFoundError):
                export_result(self.data, self.solution, path)
            shutil.copy2(cfg.RESULT_XLSX, path)
            original = path.read_bytes()
            damaged = deepcopy(self.solution)
            damaged.grid_purchase_kwh[0] += 10
            with self.assertRaises(ValueError):
                export_result(self.data, damaged, path)
            self.assertEqual(path.read_bytes(), original)
            with patch("src.question1.exporter.os.replace", side_effect=PermissionError("locked")):
                with self.assertRaisesRegex(RuntimeError, "locked"):
                    export_result(self.data, self.solution, path)
            self.assertEqual(path.read_bytes(), original)
            self.assertEqual(list(Path(folder).glob("*.tmp.xlsx")), [])
            workbook = load_workbook(path)
            workbook.active.title = "wrong"
            workbook.save(path)
            workbook.close()
            original = path.read_bytes()
            with self.assertRaisesRegex(RuntimeError, "expected sheets"):
                export_result(self.data, self.solution, path)
            self.assertEqual(path.read_bytes(), original)

    def test_summary_uses_template_labels(self):
        workbook = load_workbook(cfg.RESULT_XLSX)
        labels = [workbook.active.cell(r, 1).value for r in range(2, 146)]
        workbook.close()
        grid = np.arange(144, dtype=float)
        summary = build_table1_summary(labels, grid, self.data.price)
        self.assertEqual(summary["10:00-10:10"], labels.index("10:00-10:10"))
        with self.assertRaises(ValueError):
            build_table1_summary(["duplicate"] * 144, grid, self.data.price)
        summary2 = build_table2_summary(grid, grid)
        self.assertEqual(summary2["0:00-4:00"]["累计充电量"], sum(range(24)))


if __name__ == "__main__":
    main()
