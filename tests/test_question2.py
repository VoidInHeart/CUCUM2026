"""问题2防泄漏、LP物理约束、事后结算及导出回归测试。"""
from dataclasses import replace
from datetime import date
from unittest import TestCase, main
from unittest.mock import patch
from types import SimpleNamespace
from pathlib import Path
from tempfile import TemporaryDirectory
import shutil

import numpy as np

from src.question2 import config as cfg
from src.question2.data_loader import load_actual, load_price
from src.question2.forecast import (
    build_weekday_load_profile, compute_recent_load_trend, forecast_load_weekday_trend,
)
from src.question2.scenario_builder import (
    SCENARIO_METHODS, build_rolling_scenarios, build_scenarios, validate_scenarios,
)
from src.question2.optimizer import solve, solve_grid_first_stage_storage_recourse
from src.question2.controller import execute_causal_day
from src.question2.evaluator import evaluate, evaluate_causal_storage, evaluate_frozen_storage
from src.question2.validator import validate_plan, validate_evaluation, validate_lp
from src.question2.emergency import merge_emergency_intervals
from src.question2.exporter import export_result, read_template_labels
from src.question2.pipeline import run_annual
from src.question2.summary import build_ablation_summary, build_summaries
from openpyxl import load_workbook


class Question2Tests(TestCase):
    @classmethod
    def setUpClass(cls):
        cls.price = load_price()
        cls.actual = load_actual(price=cls.price)
        cls.scenarios = build_rolling_scenarios(cls.actual, 31)
        cls.plan = solve(cls.price.price, cls.scenarios)

    def test_attachment2_shape_and_net_load_conversion(self):
        self.assertEqual(self.actual.load_kw.shape, (365, 144))
        self.assertEqual(self.actual.pv_kw.shape, (365, 144))
        np.testing.assert_allclose(self.actual.net_load_kwh, (self.actual.load_kw - self.actual.pv_kw) / 6)

    def test_scenario_dates_and_probabilities(self):
        for index in (31, 171, 364):
            scenarios = build_rolling_scenarios(self.actual, index)
            self.assertEqual(scenarios.history_dates, self.actual.dates[index - 31:index])
            self.assertTrue(all(day < scenarios.target_date for day in scenarios.history_dates))
            self.assertEqual(scenarios.net_load_kwh.shape, (31, 144))
            self.assertAlmostEqual(scenarios.probability.sum(), 1)
        self.assertEqual(self.scenarios.history_dates[0], date(2025, 1, 1))
        self.assertEqual(build_rolling_scenarios(self.actual, 364).history_dates[-1], date(2025, 12, 30))
        for index in (0, 30, 365):
            with self.assertRaises(ValueError):
                build_rolling_scenarios(self.actual, index)

    def test_future_actuals_do_not_change_plan(self):
        index = 171
        original = build_rolling_scenarios(self.actual, index)
        net = self.actual.net_load_kwh.copy()
        net[index:] += 1e6
        modified = build_rolling_scenarios(replace(self.actual, net_load_kwh=net), index)
        np.testing.assert_array_equal(original.net_load_kwh, modified.net_load_kwh)
        first, second = solve(self.price.price, original), solve(self.price.price, modified)
        np.testing.assert_allclose(first.grid_kwh, second.grid_kwh)
        before = first.grid_kwh.copy()
        settlement = evaluate(self.price.price, first, net[index])
        self.assertGreater(settlement.emergency_total, 1e6)
        np.testing.assert_array_equal(first.grid_kwh, before)
        with self.assertRaises(ValueError):
            first.grid_kwh[0] = 0

    def test_lookahead_and_nonuniform_weights_rejected(self):
        with self.assertRaisesRegex(ValueError, "no lookahead"):
            validate_scenarios(replace(self.scenarios, history_dates=self.scenarios.history_dates[1:] + (self.scenarios.target_date,)))
        weights = np.arange(1, 32, dtype=float)
        weights /= weights.sum()
        with self.assertRaisesRegex(ValueError, "equal"):
            validate_scenarios(replace(self.scenarios, probability=weights))

    def test_corrected_and_residual_scenarios_have_no_future_leakage(self):
        index = 171
        load, pv = self.actual.load_kw.copy(), self.actual.pv_kw.copy()
        load[index:] += 1e6
        pv[index:] += 2e6
        changed = replace(self.actual, load_kw=load, pv_kw=pv,
                          net_load_kwh=(load - pv) * cfg.TIME_STEP_HOURS)
        for method in SCENARIO_METHODS:
            with self.subTest(method=method):
                first = build_scenarios(self.actual, index, method)
                second = build_scenarios(changed, index, method)
                np.testing.assert_allclose(first.load_kw, second.load_kw, rtol=0, atol=1e-10)
                np.testing.assert_allclose(first.pv_kw, second.pv_kw, rtol=0, atol=1e-10)
                np.testing.assert_allclose(first.net_load_kwh, second.net_load_kwh, rtol=0, atol=1e-10)
                np.testing.assert_array_equal(first.probability, second.probability)
                solve_first, solve_second = solve(self.price.price, first), solve(self.price.price, second)
                np.testing.assert_allclose(solve_first.grid_kwh, solve_second.grid_kwh,
                                           rtol=0, atol=cfg.RESIDUAL_TOL)

    def test_weekday_profile_and_trend_ignore_target_and_future(self):
        index = 171
        load = self.actual.load_kw.copy()
        load[index:] += 1e6
        changed = replace(self.actual, load_kw=load)
        functions = (build_weekday_load_profile, compute_recent_load_trend, forecast_load_weekday_trend)
        for function in functions:
            with self.subTest(function=function.__name__):
                np.testing.assert_allclose(function(self.actual, index), function(changed, index), rtol=0, atol=1e-10)

    def test_weighted_scenarios_are_normalized_and_nonuniform(self):
        for method in ("weekday_trend_weighted", "residual_weighted"):
            scenarios = build_scenarios(self.actual, 171, method)
            validate_scenarios(scenarios)
            self.assertAlmostEqual(scenarios.probability.sum(), 1)
            self.assertGreater(np.ptp(scenarios.probability), 0)
            self.assertTrue(all(day < scenarios.target_date for day in scenarios.history_dates))

    def test_grid_first_stage_recourse_is_feasible_and_more_flexible(self):
        scenarios = build_scenarios(self.actual, 171, "residual_weighted")
        shared = solve(self.price.price, scenarios)
        recourse = solve_grid_first_stage_storage_recourse(self.price.price, scenarios)
        self.assertLessEqual(recourse.expected_objective, shared.expected_objective + cfg.RESIDUAL_TOL)
        self.assertLessEqual(validate_plan(recourse), cfg.RESIDUAL_TOL)

    def test_grid_first_stage_recourse_plan_has_no_future_leakage(self):
        index = 171
        original = build_scenarios(self.actual, index, "residual_weighted")
        load, pv = self.actual.load_kw.copy(), self.actual.pv_kw.copy()
        load[index:] += 1e6
        pv[index:] += 2e6
        changed_data = replace(self.actual, load_kw=load, pv_kw=pv,
                               net_load_kwh=(load - pv) * cfg.TIME_STEP_HOURS)
        changed = build_scenarios(changed_data, index, "residual_weighted")
        first = solve_grid_first_stage_storage_recourse(self.price.price, original)
        second = solve_grid_first_stage_storage_recourse(self.price.price, changed)
        np.testing.assert_allclose(first.grid_kwh, second.grid_kwh, rtol=0, atol=cfg.RESIDUAL_TOL)

    def test_lp_feasible_and_terminal_soc(self):
        self.assertLessEqual(validate_plan(self.plan), 1e-5)
        self.assertAlmostEqual(self.plan.soc_end_kwh[-1], 6000, places=5)
        net = self.scenarios.net_load_kwh
        scheduled = self.plan.grid_kwh + self.plan.discharge_kwh - self.plan.charge_kwh
        emergency, surplus = np.maximum(net - scheduled, 0), np.maximum(scheduled - net, 0)
        validate_lp(self.price.price, self.scenarios, self.plan, emergency, surplus)
        emergency[0, 0] += 1
        with self.assertRaisesRegex(ValueError, "balance"):
            validate_lp(self.price.price, self.scenarios, self.plan, emergency, surplus)

    def test_emergency_multiplier_and_actual_balance(self):
        scheduled = self.plan.grid_kwh + self.plan.discharge_kwh - self.plan.charge_kwh
        actual = scheduled + np.where(np.arange(144) % 2 == 0, 100, -50)
        result = evaluate(self.price.price, self.plan, actual)
        self.assertTrue((result.emergency_kwh >= 0).all())
        self.assertAlmostEqual(result.emergency_total, 7200)
        self.assertAlmostEqual(result.emergency_cost, 500 * self.price.price[::2].sum())
        np.testing.assert_allclose(scheduled + result.emergency_kwh - result.surplus_kwh, actual)
        with self.assertRaises(ValueError):
            validate_evaluation(self.price.price, replace(result, total_cost=result.total_cost + 1))

    def test_causal_storage_keeps_grid_contract_and_is_feasible(self):
        actual = self.actual.net_load_kwh[31]
        result = evaluate_causal_storage(self.price.price, self.plan, self.scenarios, actual)
        np.testing.assert_array_equal(result.plan.grid_kwh, self.plan.grid_kwh)
        self.assertLessEqual(validate_plan(result.plan), cfg.RESIDUAL_TOL)
        validate_evaluation(self.price.price, result)

    def test_causal_storage_does_not_see_future_actual_slots(self):
        actual = self.actual.net_load_kwh[31].copy()
        changed = actual.copy()
        cutoff = 47
        changed[cutoff + 1:] += 1e6
        first = execute_causal_day(self.price.price, self.plan, self.scenarios, actual)
        second = execute_causal_day(self.price.price, self.plan, self.scenarios, changed)
        for field in ("charge_kwh", "discharge_kwh", "soc_end_kwh"):
            np.testing.assert_allclose(getattr(first, field)[:cutoff + 1], getattr(second, field)[:cutoff + 1],
                                       rtol=0, atol=cfg.RESIDUAL_TOL)

    def test_baseline_evaluator_alias_is_exact(self):
        actual = self.actual.net_load_kwh[31]
        first = evaluate(self.price.price, self.plan, actual)
        second = evaluate_frozen_storage(self.price.price, self.plan, actual)
        self.assertEqual(first.total_cost, second.total_cost)
        np.testing.assert_array_equal(first.emergency_kwh, second.emergency_kwh)

    def test_solver_failure_has_date_status_message(self):
        with patch("src.question2.optimizer.linprog", return_value=SimpleNamespace(success=False, status=2, message="infeasible")):
            with self.assertRaisesRegex(RuntimeError, "2025-02-01.*status 2.*infeasible"):
                solve(self.price.price, self.scenarios)

    def test_constant_scenario_known_cost(self):
        load = np.full((31, 144), 600.0)
        scenarios = replace(self.scenarios, net_load_kwh=np.full((31, 144), 100.0),
                            load_kw=load, pv_kw=np.zeros_like(load))
        price = np.ones(144)
        plan = solve(price, scenarios)
        self.assertAlmostEqual(plan.expected_objective, 14400, places=5)
        self.assertAlmostEqual(plan.expected_emergency_kwh, 0, places=5)

    def test_input_errors(self):
        import pandas as pd
        frame = pd.read_excel(cfg.ACTUAL_XLSX, sheet_name="小区负载")
        for case in ("shape", "date", "nan", "negative"):
            damaged = frame.copy()
            if case == "shape":
                damaged = damaged.iloc[:-1]
            elif case == "date":
                damaged.iloc[1, 0] = damaged.iloc[0, 0]
            elif case == "nan":
                damaged.iloc[3, 2] = np.nan
            else:
                damaged.iloc[3, 2] = -1
            with self.subTest(case=case), patch("src.question2.data_loader._read", return_value=damaged):
                with self.assertRaises(ValueError):
                    load_actual()

    def test_merge_emergency_intervals_and_midnight(self):
        labels = read_template_labels()
        for active, expected in (([5], [(5, 6)]), ([5, 6, 7], [(5, 8)]),
                                  ([0, 2, 3], [(0, 1), (2, 4)]), ([142, 143], [(142, 144)])):
            amounts = np.zeros(144)
            amounts[active] = 10
            events = merge_emergency_intervals(labels, amounts)
            self.assertEqual([(event.start_slot, event.end_slot) for event in events], expected)
            self.assertEqual(sum(event.amount_kwh for event in events), 10 * len(active))
            self.assertEqual(events[-1].label, labels[expected[-1][0]].split("-")[0] + "-" + labels[expected[-1][1] - 1].split("-")[1])
        self.assertEqual(events[-1].label, "23:50-0:10+1")
        amounts[:] = 0
        amounts[-1] = 2
        self.assertEqual(merge_emergency_intervals(labels, amounts)[0].label, "0:00-0:10+1")
        self.assertEqual(merge_emergency_intervals(labels, np.full(144, cfg.EMERGENCY_TOL)), [])

    def test_main_stops_before_export_on_daily_failure(self):
        from src.question2.main import main as run_main
        with patch("sys.argv", ["run_question2.py", "--no-plots"]), \
             patch("src.question2.main.run_annual", side_effect=RuntimeError("2025-06-21 solver failed")), \
             patch("src.question2.main.export_result") as export:
            self.assertEqual(run_main(), 1)
            export.assert_not_called()


class Question2ExportTests(TestCase):
    @classmethod
    def setUpClass(cls):
        cls.price = load_price()
        cls.actual = load_actual(price=cls.price)
        with patch("src.question2.pipeline.logging.info"):
            cls.results = run_annual(cls.price, cls.actual)
        cls.labels = read_template_labels()

    def test_all_334_days_complete(self):
        self.assertEqual(tuple(item.plan.date for item in self.results), cfg.OUTPUT_DATES)
        for item in self.results:
            validate_evaluation(self.price.price, item)

    def test_result2_template_and_exact_totals(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "result2.xlsx"
            shutil.copy2(cfg.RESULT_XLSX, path)
            before = load_workbook(path)
            headers = list(before["计划购电量"].values)[0]
            before.close()
            export_result(self.price.price, self.results, path)
            workbook = load_workbook(path)
            self.assertEqual(workbook.sheetnames, ["计划购电量", "充放电量", "紧急购电量"])
            grid, battery, emergency = (workbook[name] for name in workbook.sheetnames)
            self.assertEqual((grid.max_row, grid.max_column), (335, 147))
            self.assertEqual(tuple(cell.value for cell in grid[1]), headers)
            self.assertEqual(battery.max_row, 334 * 6 + 1)
            self.assertEqual(battery.max_column, 6)
            expected_events = []
            for i, item in enumerate(self.results):
                self.assertEqual(grid.cell(i + 2, 1).value.date(), item.plan.date)
                np.testing.assert_allclose([grid.cell(i + 2, t + 2).value for t in range(144)], item.plan.grid_kwh)
                self.assertAlmostEqual(grid.cell(i + 2, 146).value, item.total_grid_purchase)
                self.assertAlmostEqual(grid.cell(i + 2, 147).value, item.total_cost)
                for k in range(6):
                    row = 2 + i * 6 + k
                    self.assertAlmostEqual(battery.cell(row, 3).value, item.plan.charge_kwh[k * 24:(k + 1) * 24].sum())
                    self.assertAlmostEqual(battery.cell(row, 4).value, item.plan.discharge_kwh[k * 24:(k + 1) * 24].sum())
                    if k == 0:
                        self.assertEqual(battery.cell(row, 1).value.date(), item.plan.date)
                    else:
                        self.assertIsNone(battery.cell(row, 1).value)
                    self.assertEqual(battery.cell(row, 6).value, 6000 if k < 2 else None)
                    self.assertEqual(battery.cell(row, 5).value, ("0:00" if k == 0 else "24:00") if k < 2 else None)
                for j, event in enumerate(merge_emergency_intervals(self.labels, item.emergency_kwh)):
                    expected_events.append((item.plan.date if j == 0 else None, event.label, event.amount_kwh))
            rows = list(emergency.values)[1:]
            self.assertEqual(len(rows), len(expected_events))
            for actual, expected in zip(rows, expected_events):
                self.assertEqual(actual[0].date() if actual[0] else None, expected[0])
                self.assertEqual(actual[1], expected[1])
                self.assertAlmostEqual(actual[2], expected[2])
            workbook.close()

    def test_incomplete_and_save_failures_leave_template_unchanged(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "result2.xlsx"
            shutil.copy2(cfg.RESULT_XLSX, path)
            original = path.read_bytes()
            with self.assertRaisesRegex(ValueError, "334"):
                export_result(self.price.price, self.results[:-1], path)
            self.assertEqual(original, path.read_bytes())
            with patch("src.question2.exporter.os.replace", side_effect=PermissionError("locked")):
                with self.assertRaisesRegex(RuntimeError, "locked"):
                    export_result(self.price.price, self.results, path)
            self.assertEqual(original, path.read_bytes())
            self.assertEqual(list(Path(directory).glob("*.tmp.xlsx")), [])
            workbook = load_workbook(path)
            workbook.active.title = "wrong"
            workbook.save(path)
            workbook.close()
            original = path.read_bytes()
            with self.assertRaisesRegex(RuntimeError, "sheet names"):
                export_result(self.price.price, self.results, path)
            self.assertEqual(original, path.read_bytes())

    def test_annual_and_paper_summary(self):
        summary = build_summaries(self.results, self.labels)
        self.assertEqual(set(summary["paper_dates"]), {day.isoformat() for day in cfg.PAPER_DATES})
        annual = summary["annual"]
        self.assertAlmostEqual(annual["total_cost"], 17609791.686546136, places=5)
        self.assertAlmostEqual(annual["total_grid_purchase"], annual["planned_grid_total"] + annual["emergency_total"])
        self.assertAlmostEqual(annual["total_cost"], annual["planned_cost"] + annual["emergency_cost"])
        for item in self.results:
            if item.plan.date in cfg.PAPER_DATES:
                table1 = summary["paper_dates"][str(item.plan.date)]["表1"]
                for label, values in table1.items():
                    slot = self.labels.index(label)
                    self.assertEqual(values["计划购电量"], item.plan.grid_kwh[slot])
                    self.assertEqual(values["实际紧急购电量"], item.emergency_kwh[slot])

    def test_ablation_summary_requires_baseline(self):
        summary = build_ablation_summary({"baseline": self.results, "causal_mpc": self.results})
        self.assertEqual(summary["selected_strategy"], "baseline")
        self.assertTrue(all(row["saving_vs_A0"] == 0 for row in summary["strategies"]))
        self.assertEqual(summary["deferred_ablations"][0]["status"], "gated_out")
        with self.assertRaises(ValueError):
            build_ablation_summary({"causal_mpc": self.results})


if __name__ == "__main__":
    main()
