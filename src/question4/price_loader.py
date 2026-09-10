"""读取附件4，严格对齐附件2日期和时间列，不接触附件1价格。"""
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from types import MappingProxyType
from collections.abc import Mapping
import re
import numpy as np
from openpyxl import load_workbook

from ..question2.types import AnnualActualData, readonly
from ..pricing import price_for_date
from . import config as cfg


def normalize_time(value):
    if isinstance(value, time):
        if value.second or value.microsecond:
            raise ValueError(f"non-minute time header: {value}")
        return f"{value.hour}:{value.minute:02d}"
    match = re.fullmatch(r"(\d{1,2}):(\d{1,2})(?::00)?(\+1)?", str(value).strip())
    if not match or int(match[1]) > 23 or int(match[2]) > 59:
        raise ValueError(f"invalid time header: {value}")
    return f"{int(match[1])}:{int(match[2]):02d}{match[3] or ''}"


@dataclass(frozen=True)
class AnnualPriceData:
    dates: tuple[date, ...]
    price_yuan_per_kwh: np.ndarray
    slot_labels: tuple
    date_to_index: Mapping[date, int] = field(init=False, repr=False)

    def __post_init__(self):
        days = tuple(self.dates)
        if days != cfg.ANNUAL_DATES:
            raise ValueError("annual prices: expected 365 consecutive dates in 2025")
        prices = readonly(self.price_yuan_per_kwh)
        if prices.shape != (365, 144) or not np.isfinite(prices).all() or (prices < 0).any():
            raise ValueError("annual prices: expected 365 x 144 finite nonnegative prices")
        labels = tuple(normalize_time(value) for value in self.slot_labels)
        expected = tuple(f"{t // 60}:{t % 60:02d}" for t in range(10, 1440, 10)) + ("0:00+1",)
        if labels != expected:
            raise ValueError("annual prices: time columns must preserve the attachment order 0:10 ... 0:00+1")
        object.__setattr__(self, "dates", days)
        object.__setattr__(self, "price_yuan_per_kwh", prices)
        object.__setattr__(self, "slot_labels", labels)
        object.__setattr__(self, "date_to_index", MappingProxyType({day: i for i, day in enumerate(days)}))

    def daily_prices(self) -> dict[date, np.ndarray]:
        return {day: get_daily_price(self, day) for day in self.dates}


def get_daily_price(annual_price: AnnualPriceData, target_date: date) -> np.ndarray:
    index = annual_price.date_to_index[target_date]
    if annual_price.dates[index] != target_date:
        raise ValueError(f"{target_date}: price date index mismatch")
    return price_for_date(annual_price.price_yuan_per_kwh[index], target_date)


def validate_alignment(prices: AnnualPriceData, actual: AnnualActualData, forecasts=None):
    if prices.dates != actual.dates:
        raise ValueError("attachment 2 / attachment 4 dates differ")
    if forecasts is not None and (tuple(forecasts) != prices.dates or
            any(forecasts[day].date != day for day in prices.dates)):
        raise ValueError("attachment 3 / attachment 4 forecast dates differ")


def load_annual_price(actual: AnnualActualData, path: Path = cfg.PRICE_XLSX,
                      actual_path: Path = cfg.ACTUAL_XLSX) -> AnnualPriceData:
    path = Path(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != ["Sheet1"] or (workbook.active.max_row, workbook.active.max_column) != (366, 145):
            raise ValueError(f"{path}: expected Sheet1, A1:EO366")
        rows = iter(workbook.active.values)
        header = next(rows)
        if header[0] != "日期\\时间":
            raise ValueError(f"{path}: invalid date header")
        records = list(rows)
        days = tuple(row[0].date() if isinstance(row[0], datetime) else row[0] for row in records)
        # Excel数值单元格必须为数值；字符串、布尔值、空白不静默转换。
        if any(not isinstance(value, (int, float)) or isinstance(value, bool) for row in records for value in row[1:]):
            raise ValueError(f"{path}: price cells must be numeric")
        prices = AnnualPriceData(days, [row[1:] for row in records], tuple(header[1:]))
    finally:
        workbook.close()
    validate_alignment(prices, actual)
    workbook = load_workbook(actual_path, read_only=True, data_only=True)
    try:
        for name in ("小区负载", "光伏发电实际功率"):
            headers = next(workbook[name].values)[1:]
            if tuple(normalize_time(value) for value in headers) != prices.slot_labels:
                raise ValueError(f"{actual_path}: {name}: time columns differ from attachment 4")
    finally:
        workbook.close()
    return prices
