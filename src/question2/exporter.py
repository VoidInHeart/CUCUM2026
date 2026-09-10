"""在全部334日验证通过后填充官方模板，原子保存。"""
from copy import copy
from datetime import datetime, time
from pathlib import Path
import os
import tempfile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ..question1.summary import BLOCK_LABELS
from . import config as cfg
from .emergency import merge_emergency_intervals, validate_labels
from .summary import build_summaries
from .types import DailyEvaluation
from .validator import validate_complete
from ..pricing import PriceInput


def _check_template(workbook, path: Path) -> list[str]:
    if workbook.sheetnames != ["计划购电量", "充放电量", "紧急购电量"]:
        raise ValueError(f"{path}: exporter: invalid official sheet names/order")
    grid = workbook["计划购电量"]
    if (grid.max_row, grid.max_column) != (335, 147):
        raise ValueError(f"{path}: 计划购电量: expected 335 rows x 147 columns")
    if grid["A1"].value != "日期\\时间" or [grid.cell(1, c).value for c in (146, 147)] != ["全天购电量", "全天购电费"]:
        raise ValueError(f"{path}: 计划购电量: unexpected headers")
    dates = tuple(cell.value.date() if isinstance(cell.value, datetime) else cell.value for cell in grid["A"][1:])
    if dates != cfg.OUTPUT_DATES:
        raise ValueError(f"{path}: 计划购电量: expected all 334 official dates")
    for name, expected in (("充放电量", ["日期", "时间段", "充电量", "放电量", "时刻", "储电量"]),
                           ("紧急购电量", ["日期", "购电时间段", "购电量"])):
        sheet = workbook[name]
        if [sheet.cell(1, c).value for c in range(1, len(expected) + 1)] != expected or sheet.merged_cells:
            raise ValueError(f"{path}: {name}: unexpected headers or merged cells")
    labels = [grid.cell(1, c).value for c in range(2, 146)]
    validate_labels(labels)
    return labels


def read_template_labels(path: Path = cfg.RESULT_XLSX) -> list[str]:
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"{path}: exporter: official template missing")
    workbook = load_workbook(path)
    try:
        return _check_template(workbook, path)
    finally:
        workbook.close()


def export_result(price: PriceInput, evaluations: list[DailyEvaluation], path: Path = cfg.RESULT_XLSX) -> dict:
    validate_complete(price, evaluations)
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"{path}: exporter: official template missing")
    workbook = None
    temporary = None
    try:
        workbook = load_workbook(path)
        labels = _check_template(workbook, path)
        summary = build_summaries(evaluations, labels)
        grid, battery, emergency = (workbook[name] for name in workbook.sheetnames)
        battery_styles = [[copy(battery.cell(r, c)._style) for c in range(1, 7)] for r in range(2, 8)]
        battery_heights = [battery.row_dimensions[r].height for r in range(2, 8)]
        emergency_styles = [copy(emergency.cell(2, c)._style) for c in range(1, 4)]
        # 清掉格式示例或上次输出，保留原工作表及标题，再按样式展开。
        battery.delete_rows(2, battery.max_row - 1)
        emergency.delete_rows(2, emergency.max_row - 1)
        event_row = 2
        for i, item in enumerate(evaluations):
            plan = item.plan
            for t, value in enumerate(plan.grid_kwh):
                grid.cell(i + 2, t + 2, float(value)).number_format = "0.000000"
            grid.cell(i + 2, 146, item.total_grid_purchase).number_format = "0.000000"
            grid.cell(i + 2, 147, item.total_cost).number_format = "0.000000"
            for block, label in enumerate(BLOCK_LABELS):
                row = 2 + 6 * i + block
                section = slice(block * 24, (block + 1) * 24)
                values = [datetime.combine(plan.date, time()) if block == 0 else None, label,
                          float(plan.charge_kwh[section].sum()), float(plan.discharge_kwh[section].sum()),
                          ("0:00" if block == 0 else "24:00") if block < 2 else None,
                          cfg.INITIAL_SOC_KWH if block == 0 else float(plan.soc_end_kwh[-1]) if block == 1 else None]
                battery.row_dimensions[row].height = battery_heights[block]
                for col, value in enumerate(values, 1):
                    cell = battery.cell(row, col, value)
                    cell._style = copy(battery_styles[block][col - 1])
                    if col in (3, 4, 6):
                        cell.number_format = "0.000000"
                    elif col == 1:
                        cell.number_format = "yyyy/m/d"
            for j, event in enumerate(merge_emergency_intervals(labels, item.emergency_kwh)):
                values = [datetime.combine(plan.date, time()) if j == 0 else None, event.label, event.amount_kwh]
                for col, value in enumerate(values, 1):
                    cell = emergency.cell(event_row, col, value)
                    cell._style = copy(emergency_styles[col - 1])
                    if col == 1:
                        cell.number_format = "yyyy/m/d"
                    elif col == 3:
                        cell.number_format = "0.000000"
                event_row += 1
        # 六位小数的数值列需足够宽，避免 Excel 显示 ####。
        for col in range(2, 148):
            grid.column_dimensions[get_column_letter(col)].width = 18
        for col, width in {"A": 15, "B": 20, "C": 18, "D": 18, "E": 12, "F": 18}.items():
            battery.column_dimensions[col].width = width
        for col, width in {"A": 15, "B": 25, "C": 18}.items():
            emergency.column_dimensions[col].width = width
        with tempfile.NamedTemporaryFile(dir=path.parent, prefix=path.stem + ".", suffix=".tmp.xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        workbook.save(temporary)
        workbook.close()
        os.replace(temporary, path)
        return summary
    except Exception as exc:
        raise RuntimeError(f"{path}: exporter: {exc}") from exc
    finally:
        if workbook is not None:
            workbook.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)
