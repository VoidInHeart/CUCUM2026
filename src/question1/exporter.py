"""只填写官方模板的指定单元格，保存成功后原子替换。"""
from pathlib import Path
import os
import tempfile

from openpyxl import load_workbook

from . import config as cfg
from .data_loader import Question1Input
from .optimizer import Question1Solution
from .summary import BLOCK_LABELS, build_table1_summary, build_table2_summary
from .validator import validate_solution


def export_result(data: Question1Input, solution: Question1Solution,
                  path: Path = cfg.RESULT_XLSX) -> tuple[dict, dict]:
    # 公开导出接口自身也校验，避免调用方绕过主程序验证。
    validate_solution(data, solution)
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"{path}: exporter: official template does not exist")
    workbook = None
    temporary = None
    try:
        workbook = load_workbook(path)
        if workbook.sheetnames != ["计划购电量", "充放电量"]:
            raise ValueError("expected sheets 计划购电量, 充放电量 in official order")
        grid = workbook["计划购电量"]
        battery = workbook["充放电量"]
        if (grid.max_row, grid.max_column) != (145, 2):
            raise ValueError("计划购电量: expected 145 rows and 2 columns")
        if [grid.cell(1, c).value for c in (1, 2)] != ["时间段", "购电量"]:
            raise ValueError("计划购电量: unexpected headers")
        if [battery.cell(1, c).value for c in range(1, 6)] != ["时间段", "充电量", "放电量", "时刻", "储电量"]:
            raise ValueError("充放电量: unexpected headers")
        if [battery.cell(r, 1).value for r in range(2, 8)] != BLOCK_LABELS:
            raise ValueError("充放电量: unexpected block labels")
        if [battery["D2"].value, battery["D3"].value] != ["0:00", "24:00"]:
            raise ValueError("充放电量: unexpected SOC time labels")
        labels = [grid.cell(r, 1).value for r in range(2, 146)]
        if any(not isinstance(label, str) or not label for label in labels):
            raise ValueError("计划购电量: missing interval label")
        table1 = build_table1_summary(labels, solution.grid_purchase_kwh, data.price)
        table2 = build_table2_summary(solution.charge_kwh, solution.discharge_kwh)

        def fill(sheet, row, col, value):
            cell = sheet.cell(row, col)
            cell.value = float(value)
            cell.number_format = "0.000000"

        for i, value in enumerate(solution.grid_purchase_kwh):
            fill(grid, i + 2, 2, value)
        for k, label in enumerate(BLOCK_LABELS):
            fill(battery, k + 2, 2, table2[label]["累计充电量"])
            fill(battery, k + 2, 3, table2[label]["累计放电量"])
        fill(battery, 2, 5, cfg.INITIAL_SOC_KWH)
        fill(battery, 3, 5, solution.soc_end_kwh[-1])
        with tempfile.NamedTemporaryFile(dir=path.parent, prefix="result1.", suffix=".tmp.xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        workbook.save(temporary)
        workbook.close()
        os.replace(temporary, path)
        return table1, table2
    except Exception as exc:
        raise RuntimeError(f"{path}: exporter: {exc}") from exc
    finally:
        if workbook is not None:
            workbook.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)
