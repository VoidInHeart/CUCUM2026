"""全部334日及三次revision验证通过后，原子填充官方四工作表。"""
from copy import copy
from datetime import datetime, time
from pathlib import Path
import os
import tempfile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ..question1.summary import BLOCK_LABELS
from ..question2.emergency import validate_labels, merge_emergency_intervals
from . import config as cfg
from .validator import validate_complete

SHEETS = ["计划购电量", "调整购电量", "充放电量", "紧急购电量"]


def _check_template(workbook):
    if workbook.sheetnames != SHEETS:
        raise ValueError("expected four official sheets in order")
    labels = []
    for name in SHEETS[:2]:
        sheet = workbook[name]
        if (sheet.max_row, sheet.max_column) != (335, 147):
            raise ValueError(f"{name}: expected 335 x 147 cells")
        if sheet["A1"].value != "日期\\时间" or [sheet.cell(1, c).value for c in (146, 147)] != ["全天购电量", "全天购电费"]:
            raise ValueError(f"{name}: invalid headers")
        days = tuple(cell.value.date() if isinstance(cell.value, datetime) else cell.value for cell in sheet["A"][1:])
        if days != cfg.OUTPUT_DATES:
            raise ValueError(f"{name}: expected 334 dates in official order")
        labels.append([sheet.cell(1, c).value for c in range(2, 146)])
        validate_labels(labels[-1])
    if labels[0] != labels[1]:
        raise ValueError("plan and adjusted sheet labels differ")
    for name, headers in ((SHEETS[2], ["日期", "时间段", "充电量", "放电量", "时刻", "储电量"]),
                           (SHEETS[3], ["日期", "购电时间段", "购电量"])):
        sheet = workbook[name]
        if [sheet.cell(1, c).value for c in range(1, len(headers) + 1)] != headers or sheet.merged_cells:
            raise ValueError(f"{name}: unexpected headers or merged cells")
    return labels[0]


def read_template_labels(path: Path = cfg.RESULT_XLSX):
    workbook = None
    try:
        workbook = load_workbook(path)
        return _check_template(workbook)
    except Exception as exc:
        raise ValueError(f"{path}: question3 exporter: {exc}") from exc
    finally:
        if workbook is not None:
            workbook.close()


def export_result(price, results, path: Path = cfg.RESULT_XLSX):
    validate_complete(price, results)
    path = Path(path)
    workbook = temporary = None
    try:
        workbook = load_workbook(path)
        labels = _check_template(workbook)
        initial, adjusted, battery, emergency = (workbook[name] for name in SHEETS)
        styles = [[copy(battery.cell(r, c)._style) for c in range(1, 7)] for r in range(2, 8)]
        heights = [battery.row_dimensions[r].height for r in range(2, 8)]
        event_styles = [copy(emergency.cell(2, c)._style) for c in range(1, 4)]
        battery.delete_rows(2, battery.max_row - 1)
        emergency.delete_rows(2, emergency.max_row - 1)
        event_row = 2
        for i, result in enumerate(results):
            s = result.schedule
            for sheet, values, total, cost in ((initial, s.initial_plan.grid_kwh, result.initial_plan_kwh, result.plan_cost),
                                              (adjusted, s.executed_grid, result.actual_grid_energy_kwh, result.total_cost)):
                for col, value in enumerate([*values, total, cost], 2):
                    sheet.cell(i + 2, col, float(value)).number_format = "0.000000"
            for block, label in enumerate(BLOCK_LABELS):
                row = 2 + 6 * i + block
                section = slice(block * cfg.BLOCK_SIZE, (block + 1) * cfg.BLOCK_SIZE)
                values = [datetime.combine(s.date, time()) if block == 0 else None, label,
                          float(s.executed_charge[section].sum()), float(s.executed_discharge[section].sum()),
                          ("0:00" if block == 0 else "24:00") if block < 2 else None,
                          cfg.INITIAL_SOC_KWH if block == 0 else float(s.executed_soc_end[-1]) if block == 1 else None]
                battery.row_dimensions[row].height = heights[block]
                for col, value in enumerate(values, 1):
                    cell = battery.cell(row, col, value)
                    cell._style = copy(styles[block][col - 1])
                    if col in (3, 4, 6): cell.number_format = "0.000000"
                    elif col == 1: cell.number_format = "yyyy/m/d"
            for j, event in enumerate(merge_emergency_intervals(labels, result.real_emergency)):
                for col, value in enumerate((datetime.combine(s.date, time()) if j == 0 else None, event.label, event.amount_kwh), 1):
                    cell = emergency.cell(event_row, col, value)
                    cell._style = copy(event_styles[col - 1])
                    if col == 1: cell.number_format = "yyyy/m/d"
                    elif col == 3: cell.number_format = "0.000000"
                event_row += 1
        for sheet in (initial, adjusted):
            for col in range(2, 148):
                sheet.column_dimensions[get_column_letter(col)].width = 18
        for col, width in {"A": 15, "B": 20, "C": 18, "D": 18, "E": 12, "F": 18}.items():
            battery.column_dimensions[col].width = width
        for col, width in {"A": 15, "B": 25, "C": 18}.items():
            emergency.column_dimensions[col].width = width
        with tempfile.NamedTemporaryFile(dir=path.parent, prefix="result3.", suffix=".tmp.xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        workbook.save(temporary)
        workbook.close()
        os.replace(temporary, path)
    except Exception as exc:
        raise RuntimeError(f"{path}: question3 exporter: {exc}") from exc
    finally:
        if workbook is not None:
            workbook.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)
